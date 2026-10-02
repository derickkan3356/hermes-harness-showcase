# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx"]
# ///
"""Write the CV template: a single-column CV with docxtpl tags.

    uv run --script cv/make_template.py [out.docx]

`out` defaults to cv/template.docx.
The template is the look; the `cv` worker fills it. Restyle in Word if you
like, keeping every `{{ }}` and `{%p %}` tag as it is: each loop tag sits in a
paragraph of its own, and docxtpl removes those paragraphs when it fills the
file. Rerunning this script overwrites a restyled template.

Single column, real heading styles, no text boxes and no layout tables, so ATS
software reads it in order and an agency can paste it into their own template.

Fonts are ones Word has on Windows and macOS, so the .docx looks the same on
the recruiter's machine. The PDF is made by LibreOffice here, which needs a
metric-compatible twin: Caladea for Cambria, in cv/fonts (`docs/cv.md`).
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
BLACK, GREY = "000000", "595959"

# Cambria, a left-aligned name and section headings in navy over a navy rule, grey dates.
LOOK = dict(font="Cambria", size=10.5, name_size=24, name_align="left", accent="1F3A5F", muted=GREY,
            heading=dict(size=11, bold=True, small_caps=False, caps=True, align="left", rule=True),
            margins=(1.4, 1.8))

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}


def build(look: dict, out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    top, side = look["margins"]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(side)
    sec.top_margin = sec.bottom_margin = Cm(top)
    width = sec.page_width - 2 * Cm(side)
    styles = doc.styles

    def set_font(st, size=None, bold=None, italic=None, color=None):
        st.font.name = look["font"]
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        # The default template binds fonts to the theme; a theme font wins over a named one in Word.
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), look["font"])
        if size:
            st.font.size = Pt(size)
        if bold is not None:
            st.font.bold = bold
        if italic is not None:
            st.font.italic = italic
        st.font.color.rgb = RGBColor.from_string(color or BLACK)

    def spacing(st, before=0, after=0, keep=False, align=None, right_tab=False):
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1.0
        pf.keep_with_next = keep
        if align:
            pf.alignment = ALIGN[align]
        if right_tab:
            pf.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)

    def rule(st, color):
        ppr = st.element.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        ppr.append(bdr)

    def style(name, base="Normal"):
        if name in [s.name for s in styles]:
            return styles[name]
        st = styles.add_style(name, 1)
        st.base_style = styles[base]
        return st

    normal = styles["Normal"]
    set_font(normal, size=look["size"])
    spacing(normal)

    st = style("Name")
    set_font(st, size=look["name_size"], bold=True, color=look["accent"])
    spacing(st, after=2, align=look["name_align"])

    st = style("Contact")
    set_font(st, size=look["size"] - 1, color=look["muted"])
    spacing(st, after=4, align=look["name_align"])

    h = look["heading"]
    st = styles["Heading 1"]
    set_font(st, size=h["size"], bold=h["bold"], italic=False, color=look["accent"])
    st.font.small_caps, st.font.all_caps = h["small_caps"], h["caps"]
    spacing(st, before=9, after=3, keep=True, align=h["align"])
    if h["rule"]:
        rule(st, look["accent"])

    st = style("Company")
    set_font(st, bold=True)
    spacing(st, before=5, keep=True, right_tab=True)

    st = style("Job Title")
    set_font(st, italic=True)
    spacing(st, keep=True, right_tab=True)

    st = style("Project")
    set_font(st, size=look["size"] - 0.5, bold=True, italic=True)
    spacing(st, before=4, after=1, keep=True, right_tab=True)

    st = styles["List Bullet"]
    set_font(st)
    spacing(st, after=1)

    def para(text, st="Normal"):
        return doc.add_paragraph(text, style=st)

    def tabbed(left, right, st, right_style=None):
        p = doc.add_paragraph(style=st)
        p.add_run(left)
        p.add_run().add_tab()
        r = p.add_run(right)
        if right_style == "plain":
            r.bold = False
        r.font.color.rgb = RGBColor.from_string(look["muted"])
        return p

    def bullets(var):
        para(f"{{%p for b in {var} %}}")
        p = para("{{ b }}", "List Bullet")
        # The numbering definition sets its own indent; direct formatting is what wins over it.
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        para("{%p endfor %}")

    def project(name, after, dates=None):
        # The name stands out; the client (or kind of project) after it does not.
        p = doc.add_paragraph(style="Project")
        p.add_run(name)
        r = p.add_run(after)
        r.bold = r.italic = False
        r.font.color.rgb = RGBColor.from_string(look["muted"])
        if dates:
            p.add_run().add_tab()
            d = p.add_run(dates)
            d.bold = False
            d.font.color.rgb = RGBColor.from_string(look["muted"])
        return p

    def labelled(label, var):
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(var)

    para("{{ name }}", "Name")
    para("{{ contact }}", "Contact")

    para("Summary", "Heading 1")
    para("{{ summary }}")

    para("Experience", "Heading 1")
    para("{%p for r in roles %}")
    tabbed("{{ r.company }}", "{{ r.location }}", "Company", right_style="plain")
    para("{%p for t in r.titles %}")
    tabbed("{{ t.title }}", "{{ t.dates }}", "Job Title")
    para("{%p endfor %}")
    bullets("r.bullets")
    para("{%p for p in r.projects %}")
    project("{{ p.heading }}", "{{ p.after }}")
    bullets("p.bullets")
    para("{%p endfor %}")
    para("{%p endfor %}")

    para("{%p if projects %}")
    para("Projects", "Heading 1")
    para("{%p for p in projects %}")
    project("{{ p.heading }}", "{{ p.after }}", "{{ p.dates }}")
    para("{%p if p.url %}")
    para("{{ p.url }}", "Contact").alignment = WD_ALIGN_PARAGRAPH.LEFT
    para("{%p endif %}")
    bullets("p.bullets")
    para("{%p endfor %}")
    para("{%p endif %}")

    para("Education", "Heading 1")
    para("{%p for e in education %}")
    tabbed("{{ e.school }}", "{{ e.dates }}", "Company", right_style="plain")
    para("{{ e.degree }}{{ e.gpa }}", "Job Title")
    para("{%p endfor %}")

    para("Skills", "Heading 1")
    para("{%p for s in skills %}")
    labelled("{{ s.label }}", "{{ s.names }}")
    para("{%p endfor %}")

    para("Additional information", "Heading 1")
    labelled("Languages", "{{ languages }}")
    labelled("Right to work", "{{ right_to_work }}; notice period {{ notice_period }}")

    doc.save(out)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "template.docx"
    build(LOOK, out)
    print(out)
